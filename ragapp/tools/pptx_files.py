"""PowerPoint tools using python-pptx; edits target existing shapes instead of rebuilding slides."""
from pptx import Presentation
from ragapp.tools.definitions import Tool
from ragapp.workspace.manager import resolve_workspace_path

def _path(u,p): return resolve_workspace_path(u,p)
def read_pptx(username,relative_path):
    path=_path(username,relative_path)
    if not path.is_file(): raise FileNotFoundError(f"File not found: {relative_path}")
    prs=Presentation(str(path)); slides=[]
    for si,slide in enumerate(prs.slides):
        shapes=[]
        for shi,s in enumerate(slide.shapes):
            item={"index":shi,"type":str(s.shape_type),"left":s.left/914400,"top":s.top/914400,"width":s.width/914400,"height":s.height/914400}
            if getattr(s,"has_text_frame",False): item["text"]=s.text; item["paragraphs"]= [{"text":p.text,"runs":[{"text":r.text,"bold":r.font.bold,"italic":r.font.italic,"size":r.font.size.pt if r.font.size else None} for r in p.runs]} for p in s.text_frame.paragraphs]
            shapes.append(item)
        slides.append({"index":si,"shapes":shapes})
    return {"path":relative_path,"slide_count":len(prs.slides),"slides":slides}
def write_pptx(username,relative_path,slides):
    path=_path(username,relative_path); path.parent.mkdir(parents=True,exist_ok=True)
    if path.exists(): raise FileExistsError(f"File already exists: {relative_path}")
    prs=Presentation(); blank=prs.slide_layouts[6]
    for spec in slides:
        slide=prs.slides.add_slide(blank)
        for shape in spec.get("shapes",[]):
            box=slide.shapes.add_textbox(shape.get("left",1)*914400,shape.get("top",1)*914400,shape.get("width",8)*914400,shape.get("height",1)*914400)
            tf=box.text_frame; tf.clear(); p=tf.paragraphs[0]
            for run in shape.get("runs",[{"text":shape.get("text","")}]):
                r=p.add_run(); r.text=run.get("text",""); r.font.bold=run.get("bold"); r.font.italic=run.get("italic")
    prs.save(path); return {"path":relative_path,"status":"created"}
def edit_pptx(username,relative_path,slide_index,shape_index,operation,run_index=None,old_text=None,new_text=None):
    path=_path(username,relative_path)
    if not path.is_file(): raise FileNotFoundError(f"File not found: {relative_path}")
    prs=Presentation(str(path)); shape=prs.slides[slide_index].shapes[shape_index]
    if not getattr(shape,"has_text_frame",False): raise ValueError("Selected shape has no text frame")
    if operation == "replace_run":
        if run_index is None or new_text is None: raise ValueError("replace_run requires run_index and new_text")
        runs=[r for p in shape.text_frame.paragraphs for r in p.runs]
        if run_index < 0 or run_index >= len(runs): raise ValueError("run_index is out of range")
        runs[run_index].text=new_text
    elif operation == "replace_text_in_runs":
        if old_text is None or new_text is None: raise ValueError("replace_text_in_runs requires old_text and new_text")
        count=0
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if old_text in run.text:
                    run.text=run.text.replace(old_text,new_text); count += 1
        if not count: raise ValueError("old_text was not found within a single run")
    else:
        raise ValueError("Unsupported PowerPoint edit operation")
    prs.save(path); return {"path":relative_path,"status":"edited"}
def delete_pptx(username,relative_path):
    path=_path(username,relative_path)
    if not path.is_file(): raise FileNotFoundError(f"File not found: {relative_path}")
    path.unlink(); return {"path":relative_path,"status":"deleted"}
def build_pptx_tools(username):
    return [
      Tool("read_pptx","Read a PowerPoint as slides and shapes, including text and positions.",{"type":"object","properties":{"relative_path":{"type":"string"}},"required":["relative_path"]},lambda relative_path:read_pptx(username,relative_path)),
      Tool(
          "write_pptx",
          "Create a real .pptx from slide/shape specifications.",
          {
              "type": "object",
              "properties": {
                  "relative_path": {"type": "string"},
                  "slides": {
                      "type": "array",
                      "items": {
                          "type": "object",
                          "properties": {
                              "shapes": {
                                  "type": "array",
                                  "items": {
                                      "type": "object",
                                      "properties": {
                                          "left": {"type": "number", "description": "Position from left edge, in inches. Default 1."},
                                          "top": {"type": "number", "description": "Position from top edge, in inches. Default 1."},
                                          "width": {"type": "number", "description": "Textbox width, in inches. Default 8."},
                                          "height": {"type": "number", "description": "Textbox height, in inches. Default 1."},
                                          "text": {"type": "string", "description": "Shorthand: plain text for the shape if 'runs' is not used."},
                                          "runs": {
                                              "type": "array",
                                              "items": {
                                                  "type": "object",
                                                  "properties": {
                                                      "text": {"type": "string"},
                                                      "bold": {"type": "boolean"},
                                                      "italic": {"type": "boolean"},
                                                  },
                                                  "required": ["text"],
                                              },
                                          },
                                      },
                                  },
                              },
                          },
                          "required": ["shapes"],
                      },
                  },
              },
              "required": ["relative_path", "slides"],
          },
          lambda relative_path,slides:write_pptx(username,relative_path,slides),
      ),
      Tool("edit_pptx","Edit text in an existing PowerPoint run so the surrounding slide objects and run formatting are retained.",{"type":"object","properties":{"relative_path":{"type":"string"},"slide_index":{"type":"integer"},"shape_index":{"type":"integer"},"operation":{"type":"string","enum":["replace_run","replace_text_in_runs"]},"run_index":{"type":"integer"},"old_text":{"type":"string"},"new_text":{"type":"string"}},"required":["relative_path","slide_index","shape_index","operation"]},lambda relative_path,slide_index,shape_index,operation,run_index=None,old_text=None,new_text=None:edit_pptx(username,relative_path,slide_index,shape_index,operation,run_index,old_text,new_text)),
      Tool("delete_pptx","Delete an existing PowerPoint file from the AI workspace.",{"type":"object","properties":{"relative_path":{"type":"string"}},"required":["relative_path"]},lambda relative_path:delete_pptx(username,relative_path)),
    ]