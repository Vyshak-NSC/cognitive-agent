"""Structured DOCX tools. Edits operate on document objects so existing formatting is retained."""
import docx
from docx.enum.text import WD_BREAK
from ragapp.tools.definitions import Tool
from ragapp.workspace.manager import resolve_workspace_path


def _path(username, p): return resolve_workspace_path(username, p)

def _paragraph_json(p, index):
    return {
        "index": index, "style": p.style.name if p.style else None,
        "alignment": str(p.alignment) if p.alignment is not None else None,
        "runs": [{"index": i, "text": r.text, "bold": r.bold, "italic": r.italic, "underline": r.underline, "style": r.style.name if r.style else None} for i, r in enumerate(p.runs)],
        "text": p.text,
    }

def read_docx(username, relative_path, offset=0, limit=100):
    path = _path(username, relative_path)
    if not path.is_file(): raise FileNotFoundError(f"File not found: {relative_path}")
    doc = docx.Document(path)
    all_paragraphs = [_paragraph_json(p, i) for i, p in enumerate(doc.paragraphs)]
    return {"path": relative_path, "paragraphs": all_paragraphs[offset:offset+limit],
            "total_paragraphs": len(all_paragraphs), "offset": offset, "limit": limit,
            "tables": [...], "sections": len(doc.sections)}
    
    
def write_docx(username, relative_path, document_spec):
    path = _path(username, relative_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists(): raise FileExistsError(f"File already exists: {relative_path}")
    doc = docx.Document()
    for item in document_spec.get("paragraphs", []):
        p = doc.add_paragraph(style=item.get("style")) if item.get("style") else doc.add_paragraph()
        for run in item.get("runs", []):
            r = p.add_run(run.get("text", "")); r.bold = run.get("bold"); r.italic = run.get("italic"); r.underline = run.get("underline")
    for table in document_spec.get("tables", []):
        rows = table.get("rows", []); cols = max((len(r) for r in rows), default=1)
        t = doc.add_table(rows=len(rows), cols=cols)
        for i, row in enumerate(rows):
            for j, value in enumerate(row): t.cell(i, j).text = str(value)
    doc.save(path)
    return {"path": relative_path, "status": "created"}

def edit_docx(username, relative_path, operation, paragraph_index=None, run_index=None, old_text=None, new_text=None, text=None):
    path = _path(username, relative_path)
    if not path.is_file(): raise FileNotFoundError(f"File not found: {relative_path}")
    doc = docx.Document(path)
    if operation == "replace_run":
        if paragraph_index is None or run_index is None or new_text is None: raise ValueError("replace_run requires paragraph_index, run_index and new_text")
        doc.paragraphs[paragraph_index].runs[run_index].text = new_text
    elif operation == "replace_text_in_runs":
        if old_text is None or new_text is None: raise ValueError("replace_text_in_runs requires old_text and new_text")
        count = 0
        for p in doc.paragraphs:
            for r in p.runs:
                if old_text in r.text: r.text = r.text.replace(old_text, new_text); count += 1
        if not count: raise ValueError("old_text was not found within a single run")
    elif operation == "append_paragraph":
        p = doc.add_paragraph(text or "")
    elif operation == "delete_paragraph":
        if paragraph_index is None: raise ValueError("delete_paragraph requires paragraph_index")
        p = doc.paragraphs[paragraph_index]; p._element.getparent().remove(p._element)
    else: raise ValueError("Unsupported DOCX edit operation")
    doc.save(path)
    return {"path": relative_path, "status": "edited"}

def delete_docx(username, relative_path):
    path = _path(username, relative_path)
    if not path.is_file(): raise FileNotFoundError(f"File not found: {relative_path}")
    path.unlink(); return {"path": relative_path, "status": "deleted"}

def _edit_docx_handler(username, **args):
    """Validate model-generated arguments before calling the DOCX implementation."""
    relative_path = args.get("relative_path")
    operation = args.get("operation")

    if not relative_path:
        return {
            "status": "error",
            "error": "Missing required argument 'relative_path'. Provide the workspace-relative path to the DOCX file.",
            "received_arguments": args,
        }
    if not operation:
        return {
            "status": "error",
            "error": "Missing required argument 'operation'.",
            "received_arguments": args,
        }

    return edit_docx(
        username,
        relative_path,
        operation,
        args.get("paragraph_index"),
        args.get("run_index"),
        args.get("old_text"),
        args.get("new_text"),
        args.get("text"),
    )


def build_docx_tools(username):
    return [
        Tool("read_docx", "Read raw paragraph/table structure from a DOCX file. Only use this when editing the file or when compiled cognition entities do not contain the needed information. Do NOT use this to answer general questions about project content when the project is already compiled — use request_cognition_context instead.", {"type":"object","properties":{"relative_path":{"type":"string"}},"required":["relative_path"]}, lambda relative_path: read_docx(username, relative_path)),
        Tool("write_docx", "Create a real .docx from a structured document specification containing paragraphs/runs and tables.", {"type":"object","properties":{"relative_path":{"type":"string"},"document_spec":{"type":"object"}},"required":["relative_path","document_spec"]}, lambda relative_path, document_spec: write_docx(username, relative_path, document_spec)),
        Tool(
            "edit_docx",
            "Edit an existing DOCX while retaining its existing Word structure and formatting where possible. "
            "ALWAYS provide relative_path and operation. "
            "replace_run requires paragraph_index, run_index, and new_text. "
            "replace_text_in_runs requires old_text and new_text. "
            "append_paragraph optionally accepts text. "
            "delete_paragraph requires paragraph_index.",
            {
                "type": "object",
                "properties": {
                    "relative_path": {"type": "string", "description": "Workspace-relative path to the existing .docx file."},
                    "operation": {"type": "string", "enum": ["replace_run", "replace_text_in_runs", "append_paragraph", "delete_paragraph"]},
                    "paragraph_index": {"type": "integer"},
                    "run_index": {"type": "integer"},
                    "old_text": {"type": "string"},
                    "new_text": {"type": "string"},
                    "text": {"type": "string"},
                },
                "required": ["relative_path", "operation"],
            },
            lambda **args: _edit_docx_handler(username, **args),
        ),
        Tool("delete_docx", "Delete an existing DOCX from the AI workspace.", {"type":"object","properties":{"relative_path":{"type":"string"}},"required":["relative_path"]}, lambda relative_path: delete_docx(username, relative_path)),
    ]
