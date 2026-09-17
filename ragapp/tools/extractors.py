"""Deterministic text extraction used only for cognition compilation and inspection."""
from pathlib import Path
import json

def extract_text(path: Path) -> str:
    ext=path.suffix.lower()
    if ext in {'.txt','.md','.py','.js','.ts','.html','.css','.json','.yaml','.yml','.csv','.sql','.xml','.log'}:
        return path.read_text(encoding='utf-8', errors='replace')
    if ext=='.pdf':
        from pypdf import PdfReader
        r=PdfReader(str(path)); return '\n'.join((p.extract_text() or '') for p in r.pages)
    if ext=='.docx':
        from docx import Document
        d=Document(str(path)); return '\n'.join([p.text for p in d.paragraphs]+[c.text for t in d.tables for row in t.rows for c in row.cells])
    if ext=='.xlsx':
        from openpyxl import load_workbook
        wb=load_workbook(path, read_only=True, data_only=True); out=[]
        for ws in wb.worksheets:
            out.append(f'[{ws.title}]')
            for row in ws.iter_rows(values_only=True): out.append(' | '.join('' if v is None else str(v) for v in row))
        return '\n'.join(out)
    if ext=='.pptx':
        from pptx import Presentation
        prs=Presentation(str(path)); return '\n'.join(sh.text for sl in prs.slides for sh in sl.shapes if hasattr(sh,'text'))
    return f"Binary artifact: {path.name}; inspect it with the artifact tools."
