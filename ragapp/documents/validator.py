"""Artifact QA: reject materially incomplete transformations."""
from __future__ import annotations
import re
from pypdf import PdfReader
from docx import Document
from pptx import Presentation

def _tokens(text): return re.findall(r"[\w'-]+",(text or '').lower(),flags=re.UNICODE)
def _coverage(model, output, min_cov):
    src=_tokens(model.source_text()); out=_tokens(output); sc={}; oc={}
    for t in src: sc[t]=sc.get(t,0)+1
    for t in out: oc[t]=oc.get(t,0)+1
    matched=sum(min(n,oc.get(t,0)) for t,n in sc.items()); cov=matched/max(1,len(src)); words=set(out)
    missing=[b.text for b in model.blocks if b.kind=='heading' and b.text and not set(_tokens(b.text)).issubset(words)]
    issues=[]
    if cov<min_cov: issues.append(f'content coverage {cov:.1%} is below required {min_cov:.0%}')
    if missing: issues.append(f'{len(missing)} source headings missing from rendered artifact')
    return cov,missing,issues

def validate_artifact(model,path,min_token_coverage=.94):
    suffix=path.suffix.lower(); meta={}
    if suffix=='.pdf':
        r=PdfReader(str(path)); text='\n'.join((p.extract_text() or '') for p in r.pages); meta['page_count']=len(r.pages)
    elif suffix=='.docx':
        d=Document(str(path)); text='\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for row in t.rows for c in row.cells); meta['section_count']=len(d.sections)
    elif suffix=='.pptx':
        p=Presentation(str(path)); text='\n'.join(sh.text for s in p.slides for sh in s.shapes if getattr(sh,'has_text_frame',False)); text+='\n'+'\n'.join(c.text for s in p.slides for sh in s.shapes if getattr(sh,'has_table',False) for row in sh.table.rows for c in row.cells); meta['slide_count']=len(p.slides)
    elif suffix in {'.html','.htm','.md','.markdown'}: text=path.read_text(encoding='utf-8')
    else: raise ValueError(f'No validator for {suffix}')
    cov,missing,issues=_coverage(model,text,min_token_coverage)
    if not path.exists() or path.stat().st_size==0: issues.append('artifact is empty')
    return {'ok':not issues,'token_coverage':round(cov,4),'missing_headings':missing[:20],'issues':issues,**meta}

def validate_pdf(model,pdf_path,min_token_coverage=.94): return validate_artifact(model,pdf_path,min_token_coverage)
