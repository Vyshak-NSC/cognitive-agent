from __future__ import annotations
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_BREAK
from .themes import get_theme

def render_docx(model, path, theme_name='professional'):
    theme=get_theme(theme_name); doc=Document(); sec=doc.sections[0]
    sec.top_margin=Pt(theme.margin_top); sec.bottom_margin=Pt(theme.margin_bottom); sec.left_margin=Pt(theme.margin_left); sec.right_margin=Pt(theme.margin_right)
    styles=doc.styles
    styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(theme.body_size)
    for n,size in [(1,theme.h1_size),(2,theme.h2_size),(3,theme.h3_size)]:
        s=styles[f'Heading {n}']; s.font.name='Aptos Display'; s.font.size=Pt(size); s.font.bold=True
    for b in model.blocks:
        if b.kind=='heading': doc.add_heading(b.text, level=min(3,b.level or 1))
        elif b.kind=='paragraph': doc.add_paragraph(b.text)
        elif b.kind=='callout':
            p=doc.add_paragraph(); r=p.add_run(b.text); r.bold=True
        elif b.kind=='list_item': doc.add_paragraph(b.text, style='List Number' if b.ordered else 'List Bullet')
        elif b.kind=='table' and b.rows:
            cols=max(len(r) for r in b.rows); t=doc.add_table(rows=0, cols=cols); t.style='Table Grid'
            for ri,row in enumerate(b.rows):
                cells=t.add_row().cells
                for i,v in enumerate(row): cells[i].text=str(v)
                if ri==0:
                    for c in cells:
                        for r in c.paragraphs[0].runs: r.bold=True
        elif b.kind=='page_break': doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.save(str(path)); return path
