"""Project file manager for Workspace and Source: breadcrumb navigation, multi-select table and bulk copy/move/delete/download."""
from pathlib import Path
import base64, hashlib, html, io, mimetypes, re, shutil, zipfile
from datetime import datetime
from io import BytesIO
import pandas as pd
import streamlit as st
from ragapp.core.project_files import ProjectFileService

TEXT_EXTENSIONS={'.txt','.md','.py','.js','.mjs','.ts','.tsx','.jsx','.json','.yaml','.yml','.csv','.xml','.html','.htm','.css','.sql','.toml','.ini','.cfg','.env','.java','.kt','.go','.rs','.c','.cpp','.h','.hpp','.sh','.bat','.ps1','.r','.tex'}
AREAS={'workspace':('Workspace','Working files and generated artifacts.'),'source':('Source','Source material used by cognition.')}

def _safe(root,rel=''):
    root=Path(root).resolve(); p=(root/rel).resolve(); p.relative_to(root); return p

def _root(store,area):
    root=store.workspace if area=='workspace' else store.source; root.mkdir(parents=True,exist_ok=True); return root

def _is_text(p): return p.suffix.lower() in TEXT_EXTENSIONS or p.name.lower() in {'dockerfile','makefile'}
def _size(n):
    for u in ('B','KB','MB','GB'):
        if n<1024 or u=='GB': return f'{n:.0f} {u}' if u=='B' else f'{n:.1f} {u}'
        n/=1024

def _folders(root):
    root=Path(root).resolve(); return ['/']+sorted((p.relative_to(root).as_posix() for p in root.rglob('*') if p.is_dir()),key=str.lower)

def _entries(root,folder):
    root=Path(root).resolve(); cur=_safe(root,'' if folder=='/' else folder)
    out=[]
    for p in cur.iterdir():
        st_=p.stat(); isdir=p.is_dir()
        out.append({'name':p.name,'path':p.relative_to(root).as_posix(),'type':'Folder' if isdir else 'File','bytes':0 if isdir else st_.st_size,'size':'' if isdir else _size(st_.st_size),'modified':datetime.fromtimestamp(st_.st_mtime).strftime('%Y-%m-%d %H:%M')})
    return sorted(out,key=lambda x:(x['type']!='Folder',x['name'].lower()))

def _clean_name(name):
    clean=Path(name).name
    if not clean or clean!=name or clean in {'.','..'}: raise ValueError('Enter one valid file or folder name.')
    return clean

def _asset_data_uri(path):
    mime,_=mimetypes.guess_type(path.name); mime=mime or 'application/octet-stream'; return f"data:{mime};base64,{base64.b64encode(path.read_bytes()).decode('ascii')}"

def _resolve_local_asset(html_file,reference,project_root):
    reference=html.unescape(reference).strip()
    if not reference or reference.startswith(('#','data:','http://','https://','//','mailto:','javascript:')): return None
    clean=reference.split('#',1)[0].split('?',1)[0]
    try: p=(html_file.parent/clean).resolve(); p.relative_to(Path(project_root).resolve())
    except (ValueError,OSError): return None
    return p if p.is_file() else None

def _prepare_html_preview(html_file,project_root):
    source=html_file.read_text(encoding='utf-8',errors='replace')
    def attr(m):
        a=_resolve_local_asset(html_file,m.group(3),project_root)
        if a is None: return m.group(0)
        try: return f'{m.group(1)}{m.group(2)}{_asset_data_uri(a)}{m.group(2)}'
        except OSError: return m.group(0)
    return re.sub(r'(\b(?:src|poster|href)\s*=\s*)([\'\"])([^\'\"]+)\2',attr,source,flags=re.I)

def _extract_zip_into(data,root,folder,overwrite=False):
    base=_safe(root,'' if folder=='/' else folder); out={'written':[],'skipped':[],'rejected':[]}
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        for info in zf.infolist():
            name=info.filename
            if not name or name.endswith('/') or name.startswith('__MACOSX/'): continue
            try: target=(base/name).resolve(); target.relative_to(Path(root).resolve())
            except (ValueError,RuntimeError): out['rejected'].append(name); continue
            if target.exists() and not overwrite: out['skipped'].append(name); continue
            target.parent.mkdir(parents=True,exist_ok=True); target.write_bytes(zf.read(info)); out['written'].append(name)
    return out

def _preview_docx(path):
    """Render Word document text and tables without treating DOCX as raw binary."""
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError:
        st.error('DOCX preview requires python-docx. Install it with: pip install python-docx')
        return

    try:
        document = Document(BytesIO(path.read_bytes()))
        rendered_any = False
        blocks = document.iter_inner_content() if hasattr(document, 'iter_inner_content') else [*document.paragraphs, *document.tables]

        for block in blocks:
            if isinstance(block, Paragraph):
                text = block.text.strip()
                if not text:
                    continue
                rendered_any = True
                style_name = (block.style.name or '') if block.style else ''
                if style_name.startswith('Heading'):
                    try:
                        level = int(style_name.split()[-1])
                    except (TypeError, ValueError):
                        level = 3
                    st.markdown(f"{'#' * max(1, min(level, 6))} {text}")
                elif style_name in {'Title', 'Subtitle'}:
                    st.markdown(f'## {text}')
                else:
                    st.write(text)
            elif isinstance(block, Table):
                rows = [[cell.text for cell in row.cells] for row in block.rows]
                if rows:
                    rendered_any = True
                    st.table(rows)

        if not rendered_any:
            st.info('This Word document contains no previewable text or tables.')
    except Exception as exc:
        st.error(f'Could not preview Word document `{path.name}`: {exc}')


def _preview(path,root):
    suffix = path.suffix.lower()
    if _is_text(path):
        if suffix in {'.html','.htm'}: st.components.v1.html(_prepare_html_preview(path,root),height=650,scrolling=True)
        elif suffix == '.md': st.markdown(path.read_text(encoding='utf-8',errors='replace'))
        else: st.code(path.read_text(encoding='utf-8',errors='replace'),language=path.suffix.lstrip('.') or 'text')
    elif suffix in {'.png','.jpg','.jpeg','.gif','.webp'}: st.image(path.read_bytes())
    elif suffix=='.pdf': st.components.v1.html(f'<iframe src="{_asset_data_uri(path)}" width="100%" height="700"></iframe>',height=710)
    elif suffix=='.docx': _preview_docx(path)
    else: st.info(f'Binary file · {_size(path.stat().st_size)}. Use Download or Replace.')

# ---------------------------------------------------------------------------
# Bulk-operation helpers (pure logic, no Streamlit)
# ---------------------------------------------------------------------------
MAX_ZIP_BYTES = 500 * 1024 * 1024

def _rel(folder):
    """UI folder ('/' = area root) -> service-relative folder ('' = area root)."""
    return '' if folder in ('', '/') else folder

def _join(folder, name):
    return name if folder in ('', '/') else f'{folder}/{name}'

def _folder_label(f):
    return '🏠 Root' if f == '/' else '\u2003' * f.count('/') + '📁 ' + f.rsplit('/', 1)[-1]

def _parent(folder):
    return '/' if '/' not in folder else folder.rsplit('/', 1)[0]

def _transfer_problems(files, src_area, rels, dst_area, dst_folder, move):
    """Return human-readable problems that would make the whole batch fail (nothing is written)."""
    problems = []
    dst_rel = _rel(dst_folder)
    for rel in rels:
        name = Path(rel).name
        src = files.path(src_area, rel)
        target_rel = _join(dst_folder, name)
        if src_area == dst_area and target_rel == rel:
            problems.append(f"'{name}' is already in that folder.")
            continue
        if src.is_dir() and src_area == dst_area and (dst_rel == rel or dst_rel.startswith(rel + '/')):
            problems.append(f"Cannot {'move' if move else 'copy'} folder '{name}' into itself.")
            continue
        if files.path(dst_area, target_rel).exists():
            problems.append(f"'{name}' already exists in the destination.")
    return problems

def _bulk_transfer(files, src_area, rels, dst_area, dst_folder, move):
    problems = _transfer_problems(files, src_area, rels, dst_area, dst_folder, move)
    if problems:
        raise ValueError('\n'.join(problems))
    fn = files.move_many if move else files.copy_many
    verb = 'Move' if move else 'Copy'
    return fn(src_area, rels, dst_area, _rel(dst_folder), description=f'{verb} {len(rels)} item(s) {src_area} -> {dst_area}/{_rel(dst_folder) or ""}'.rstrip('/'))

def _items_size(root, rels):
    total = 0
    for rel in rels:
        p = _safe(root, rel)
        total += sum(f.stat().st_size for f in p.rglob('*') if f.is_file()) if p.is_dir() else p.stat().st_size
    return total

def _zip_items(root, rels, base):
    """Zip files/folders; archive paths are relative to `base` (the folder being browsed)."""
    root = Path(root).resolve(); base_p = _safe(root, '' if base == '/' else base)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as zf:
        for rel in rels:
            p = _safe(root, rel)
            if p.is_dir():
                found = False
                for f in sorted(p.rglob('*')):
                    if f.is_file():
                        zf.write(f, f.relative_to(base_p).as_posix()); found = True
                if not found:
                    zf.writestr(p.relative_to(base_p).as_posix() + '/', '')
            else:
                zf.write(p, p.relative_to(base_p).as_posix())
    return buf.getvalue()

# ---------------------------------------------------------------------------
# UI state helpers
# ---------------------------------------------------------------------------
def _bump(area, select_all=False):
    st.session_state[f'fm_nonce_{area}'] = st.session_state.get(f'fm_nonce_{area}', 0) + 1
    st.session_state[f'fm_selall_{area}'] = select_all

def _goto(area, folder):
    st.session_state[f'fm_cwd_{area}'] = folder
    _bump(area)

def _done(area, message):
    _bump(area)
    st.session_state['fm_flash'] = message
    st.rerun()

def _fmt_selection(sel):
    nf = sum(1 for e in sel if e['type'] == 'Folder'); nfi = len(sel) - nf
    parts = ([f'{nf} folder{"s" if nf != 1 else ""}'] if nf else []) + ([f'{nfi} file{"s" if nfi != 1 else ""}'] if nfi else [])
    return f"{len(sel)} selected · {', '.join(parts)}"

def _render_breadcrumbs(area, cwd, subfolders):
    parts = [] if cwd == '/' else cwd.split('/')
    crumbs = [('🏠 Root', '/')] + [(seg, '/'.join(parts[:i + 1])) for i, seg in enumerate(parts)]
    widths = ([0.6] if cwd != '/' else []) + [1] * len(crumbs) + [max(0.5, 6 - len(crumbs))]
    cols = st.columns(widths, gap='small'); ci = 0
    if cwd != '/':
        cols[ci].button('⬆ Up', key=f'fm_up_{area}', on_click=_goto, args=(area, _parent(cwd)), help='Go to parent folder'); ci += 1
    for i, (label, path) in enumerate(crumbs):
        last = i == len(crumbs) - 1
        cols[ci].button(label, key=f'fm_crumb_{area}_{i}', on_click=_goto, args=(area, path), disabled=last, type='primary' if last else 'secondary'); ci += 1
    if subfolders:
        st.caption('Folders here — click to open')
        shown = subfolders[:24]
        for start in range(0, len(shown), 6):
            row = st.columns(6, gap='small')
            for c, e in zip(row, shown[start:start + 6]):
                c.button(f"📁 {e['name']}", key=f"fm_sub_{area}_{e['path']}", on_click=_goto, args=(area, e['path']))
        if len(subfolders) > len(shown):
            st.caption(f'+{len(subfolders) - len(shown)} more — tick a folder in the table and press Open.')

# ---------------------------------------------------------------------------
# Selection toolbar + detail panel
# ---------------------------------------------------------------------------
def _transfer_popover(files, store, area, rels, move, nonce, folders_by_area):
    verb = 'Move' if move else 'Copy'; tag = verb.lower()
    with st.popover(f'{verb} to…'):
        dst_area = st.radio('Destination area', list(AREAS), index=list(AREAS).index(area), format_func=lambda a: AREAS[a][0], horizontal=True, key=f'fm_{tag}_area_{area}_{nonce}')
        dst_folder = st.selectbox('Destination folder', folders_by_area[dst_area], format_func=_folder_label, key=f'fm_{tag}_dst_{area}_{nonce}_{dst_area}')
        if st.button(f'{verb} {len(rels)} item(s) here', type='primary', key=f'fm_{tag}_go_{area}_{nonce}'):
            try:
                _bulk_transfer(files, area, rels, dst_area, dst_folder, move)
            except Exception as ex:
                st.error(str(ex)); return
            _done(area, f"{'Moved' if move else 'Copied'} {len(rels)} item(s) to {AREAS[dst_area][0]}{'' if dst_folder == '/' else ' / ' + dst_folder}.")

def _render_selection_panel(store, files, area, cwd, root, sel, folders_by_area):
    nonce = st.session_state.get(f'fm_nonce_{area}', 0)
    rels = [e['path'] for e in sel]
    st.markdown(f'**{_fmt_selection(sel)}**')
    c_open, c_copy, c_move, c_ren, c_dl, c_del = st.columns(6, gap='small')

    with c_open:
        one_folder = len(sel) == 1 and sel[0]['type'] == 'Folder'
        st.button('Open', key=f'fm_open_{area}_{nonce}', disabled=not one_folder, on_click=_goto, args=(area, sel[0]['path']) if one_folder else (area, cwd))
    with c_copy:
        _transfer_popover(files, store, area, rels, False, nonce, folders_by_area)
    with c_move:
        _transfer_popover(files, store, area, rels, True, nonce, folders_by_area)
    with c_ren:
        with st.popover('Rename', disabled=len(sel) != 1):
            if len(sel) == 1:
                new_name = st.text_input('New name', value=sel[0]['name'], key=f'fm_rename_{area}_{nonce}_{sel[0]["path"]}')
                if st.button('Apply', type='primary', key=f'fm_rename_go_{area}_{nonce}'):
                    try:
                        name = _clean_name(new_name.strip()); target = _join(cwd, name)
                        if target != sel[0]['path'] and files.path(area, target).exists(): raise FileExistsError(f"'{name}' already exists.")
                        if target != sel[0]['path']: files.move(area, sel[0]['path'], area, target)
                    except Exception as ex:
                        st.error(str(ex))
                    else:
                        _done(area, f"Renamed to '{name}'.")
    with c_dl:
        with st.popover('Download'):
            if len(sel) == 1 and sel[0]['type'] == 'File':
                p = _safe(root, sel[0]['path'])
                st.download_button('Download file', p.read_bytes(), file_name=p.name, key=f'fm_dl_{area}_{nonce}')
            else:
                total = _items_size(root, rels)
                if total > MAX_ZIP_BYTES:
                    st.error(f'Selection is {_size(total)}; ZIP downloads are limited to {_size(MAX_ZIP_BYTES)}.')
                else:
                    zkey = 'fm_zip_' + hashlib.md5('|'.join(rels).encode()).hexdigest()
                    st.caption(f'{_size(total)} to compress')
                    if st.button('Build ZIP', key=f'fm_zipbuild_{area}_{nonce}'):
                        st.session_state[zkey] = _zip_items(root, rels, cwd)
                    if st.session_state.get(zkey):
                        st.download_button('Download ZIP', st.session_state[zkey], file_name=f'{area}_selection.zip', mime='application/zip', key=f'fm_zipdl_{area}_{nonce}')
    with c_del:
        with st.popover('Delete'):
            st.warning(f'Permanently delete {len(sel)} item(s)? This is recorded in git history and can be restored.')
            st.caption(', '.join(e['name'] for e in sel[:8]) + (f' … +{len(sel) - 8} more' if len(sel) > 8 else ''))
            if st.button(f'Delete {len(sel)} item(s)', type='primary', key=f'fm_del_go_{area}_{nonce}'):
                try:
                    files.delete_many(area, rels, description=f'Delete {len(rels)} item(s) from {area}')
                except Exception as ex:
                    st.error(str(ex)); return
                _done(area, f'Deleted {len(rels)} item(s).')

    if len(sel) == 1 and sel[0]['type'] == 'File':
        e = sel[0]; p = _safe(root, e['path'])
        tab_view, tab_edit = st.tabs(['Preview', 'Edit / replace'])
        with tab_view:
            _preview(p, root)
        with tab_edit:
            if _is_text(p):
                text = st.text_area('Contents', p.read_text(encoding='utf-8', errors='replace'), height=430, key=f'edit_{area}_{e["path"]}')
                if st.button('Save changes', type='primary', key=f'save_{area}_{e["path"]}'):
                    files.write_text(area, e['path'], text, overwrite=True, description=f'Edit {area}/{e["path"]}')
                    st.success('Saved.')
            else:
                replacement = st.file_uploader('Replace binary file', key=f'replace_{area}_{e["path"]}')
                if replacement and st.button('Replace', type='primary', key=f'replace_go_{area}_{e["path"]}'):
                    files.write_bytes(area, e['path'], replacement.getvalue(), overwrite=True, description=f'Replace {area}/{e["path"]}')
                    st.rerun()

# ---------------------------------------------------------------------------
# Main entry point
# ---------------------------------------------------------------------------
def render_file_manager(store):
    st.subheader('Files')
    st.caption('Browse with the breadcrumb, tick any number of items, then copy, move, delete or download them in one go.')
    files = ProjectFileService(store)
    flash = st.session_state.pop('fm_flash', None)
    if flash: st.success(flash)

    st.session_state.setdefault('fm_area', 'workspace')
    area = st.segmented_control('Area', list(AREAS), format_func=lambda x: AREAS[x][0], key='fm_area') or 'workspace'
    root = _root(store, area).resolve()

    ck = f'fm_cwd_{area}'; cwd = st.session_state.get(ck, '/')
    while cwd != '/' and not (root / cwd).is_dir(): cwd = _parent(cwd)
    st.session_state[ck] = cwd
    nonce = st.session_state.get(f'fm_nonce_{area}', 0)

    all_entries = _entries(root, cwd)
    subfolders = [e for e in all_entries if e['type'] == 'Folder']
    _render_breadcrumbs(area, cwd, subfolders)

    with st.popover('＋ Add / import'):
        tab1, tab2 = st.tabs(['Create', 'Upload'])
        with tab1:
            kind = st.radio('Create', ['File', 'Folder'], horizontal=True, key=f'new_kind_{area}'); name = st.text_input('Name', key=f'new_name_{area}')
            content = st.text_area('Initial text / content', height=140, key=f'new_content_{area}') if kind == 'File' else ''
            if st.button('Create', type='primary', key=f'create_{area}'):
                try:
                    target = _safe(root, _join(cwd, _clean_name(name)))
                    if target.exists(): raise FileExistsError('Destination already exists.')
                    rel = target.relative_to(root).as_posix()
                    if kind == 'Folder': files.create_folder(area, rel)
                    else: files.write_text(area, rel, content)
                    _bump(area); st.rerun()
                except Exception as e: st.error(str(e))
        with tab2:
            ups = st.file_uploader('Files', accept_multiple_files=True, key=f'ups_{area}_{nonce}')
            if ups and st.button('Upload', type='primary', key=f'upload_{area}'):
                try:
                    batch = [(_join(cwd, Path(up.name).name), up.getvalue()) for up in ups]
                    files.write_many(area, batch, description=f'Upload {len(batch)} file(s) to {area}')
                    _bump(area); st.rerun()
                except Exception as e: st.error(str(e))
            z = st.file_uploader('Or extract a ZIP', type=['zip'], key=f'zip_{area}_{nonce}'); overwrite = st.checkbox('Overwrite existing', key=f'ow_{area}')
            if z and st.button('Extract ZIP', key=f'extract_{area}'):
                try:
                    files.vcs.checkpoint(f'Pre-change: extract ZIP into {area}/{cwd}')
                    _extract_zip_into(z.getvalue(), root, cwd, overwrite)
                    files.vcs.commit(f'Extract ZIP into {area}/{cwd}', bind_timeline=False)
                    _bump(area); st.rerun()
                except Exception as e: st.error(str(e))

    if not all_entries:
        st.info('This folder is empty.'); return

    q = st.text_input('Filter', placeholder='Filter this folder by name…', key=f'fm_filter_{area}_{cwd}', label_visibility='collapsed').strip().lower()
    entries = [e for e in all_entries if q in e['name'].lower()] if q else all_entries
    if not entries:
        st.info('No items match the filter.'); return

    b1, b2, b3 = st.columns([1, 1, 6], gap='small')
    b1.button('Select all', key=f'fm_selall_btn_{area}', on_click=_bump, args=(area, True))
    b2.button('Clear', key=f'fm_clear_btn_{area}', on_click=_bump, args=(area, False))
    b3.caption(f'{len(entries)} item(s) in {"root" if cwd == "/" else cwd}')

    default = bool(st.session_state.get(f'fm_selall_{area}', False))
    df = pd.DataFrame({
        'Select': [default] * len(entries),
        'Name': [('📁 ' if e['type'] == 'Folder' else '📄 ') + e['name'] for e in entries],
        'Type': [e['type'] for e in entries],
        'Size': [e['size'] for e in entries],
        'Modified': [e['modified'] for e in entries],
    })
    edited = st.data_editor(
        df, hide_index=True, disabled=['Name', 'Type', 'Size', 'Modified'],
        column_config={'Select': st.column_config.CheckboxColumn('✓', width='small')},
        key=f'fm_tbl_{area}_{cwd}_{nonce}_{q}',
    )
    sel = [entries[i] for i in edited.index[edited['Select']]]
    if not sel:
        st.caption('Tick one or more rows to enable Copy / Move / Rename / Download / Delete.')
        return
    folders_by_area = {a: _folders(_root(store, a)) for a in AREAS}
    _render_selection_panel(store, files, area, cwd, root, sel, folders_by_area)