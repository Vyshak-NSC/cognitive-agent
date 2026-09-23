"""Format-aware structural document parsing for cognition compilation.

The parser keeps rich structure in memory while compiling, but the persisted
cognition/document index stores only compact locators and relationships --
never the source document body.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
import hashlib
import html
import re
from typing import Any


@dataclass
class DocumentNode:
    id: str
    type: str
    order: int
    text: str = ""
    title: str | None = None
    parent_id: str | None = None
    locator: dict[str, Any] = field(default_factory=dict)
    attributes: dict[str, Any] = field(default_factory=dict)
    children: list[str] = field(default_factory=list)


@dataclass
class DocumentSegment:
    id: str
    title: str
    node_ids: list[str]
    section_ids: list[str]
    text: str
    locator: dict[str, Any]
    chapter_id: str | None = None
    scene_id: str | None = None


@dataclass
class ParsedDocument:
    artifact_id: str
    area: str
    path: str
    format: str
    title: str
    nodes: dict[str, DocumentNode]
    root_id: str
    segments: list[DocumentSegment]
    metadata: dict[str, Any] = field(default_factory=dict)

    @property
    def node_count(self) -> int:
        return len(self.nodes)


def _slug(value: str) -> str:
    value = re.sub(r"[^A-Za-z0-9_-]+", "_", str(value).strip().lower())
    return value.strip("_") or "node"


def _sha(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8", errors="replace")).hexdigest()[:16]


_CHAPTER_RE = re.compile(
    r"^\s*(?:chapter|chap\.?|part)\s+([0-9ivxlcdm]+)\b(?:\s*[:.\-–—]\s*)?(.*)$",
    re.I,
)
_SCENE_RE = re.compile(r"^\s*(?:scene|act)\s+([0-9ivxlcdm]+)\b(?:\s*[:.\-–—]\s*)?(.*)$", re.I)
_NUMBERED_HEADING_RE = re.compile(r"^\s*(\d+(?:\.\d+)*)(?:[.)]|\s+)\s*(.+?)\s*$")


def _heading_kind(text: str, style: str | None = None) -> tuple[str | None, str | None]:
    t = " ".join(str(text or "").split())
    if not t:
        return None, None
    if style:
        s = style.lower()
        if "title" in s:
            return "title", t
        if "heading 1" in s or s.endswith("heading 1"):
            return "chapter", t
        if "heading 2" in s or s.endswith("heading 2"):
            return "section", t
        if "heading 3" in s or s.endswith("heading 3"):
            return "subsection", t
    m = _CHAPTER_RE.match(t)
    if m:
        return "chapter", t
    m = _SCENE_RE.match(t)
    if m:
        return "scene", t
    m = _NUMBERED_HEADING_RE.match(t)
    if m:
        number = m.group(1)
        # Enterprise documents commonly use 1., 1.1, 1.1.1 style headings.
        # Treat top-level numbers as sections and deeper numbers as subsections.
        return ("section" if number.count(".") == 0 else "subsection"), t
    if len(t) <= 100 and (t.isupper() or t.endswith(":")):
        return "heading", t
    return None, None


def _new_node(nodes, node_type, order, text="", title=None, parent_id=None, locator=None, attributes=None):
    nid = f"node:{order}:{_sha(f'{node_type}|{order}|{title or text[:80]}')}"
    while nid in nodes:
        nid += "x"
    node = DocumentNode(
        id=nid,
        type=node_type,
        order=order,
        text=text,
        title=title,
        parent_id=parent_id,
        locator=locator or {},
        attributes=attributes or {},
    )
    nodes[nid] = node
    if parent_id and parent_id in nodes:
        nodes[parent_id].children.append(nid)
    return node


def _sectionize(nodes: dict[str, DocumentNode], root_id: str, raw_nodes: list[DocumentNode], *, max_chars: int) -> list[DocumentSegment]:
    """Turn ordered content nodes into structure-aware LLM segments."""
    current_chapter = None
    current_scene = None
    current_section = None
    sections: list[DocumentNode] = []

    for node in raw_nodes:
        if node.type == "chapter":
            current_chapter = node
            current_scene = None
            current_section = node
            sections.append(node)
        elif node.type == "scene":
            current_scene = node
            current_section = node
            sections.append(node)
        elif node.type in {"section", "subsection", "heading"}:
            current_section = node
            sections.append(node)
        elif node.text:
            node.attributes.setdefault("chapter_id", current_chapter.id if current_chapter else None)
            node.attributes.setdefault("scene_id", current_scene.id if current_scene else None)
            node.attributes.setdefault("section_id", current_section.id if current_section else None)

    # Build compact segments. We do not persist text later; it is only the LLM input.
    segments: list[DocumentSegment] = []
    by_context: dict[tuple[str | None, str | None], list[DocumentNode]] = {}
    for node in raw_nodes:
        if not node.text:
            continue
        key = (node.attributes.get("chapter_id"), node.attributes.get("scene_id"))
        by_context.setdefault(key, []).append(node)

    for (chapter_id, scene_id), items in by_context.items():
        buf: list[DocumentNode] = []
        size = 0
        part = 0
        for node in items:
            extra = len(node.text) + (2 if buf else 0)
            if buf and size + extra > max_chars:
                part += 1
                segments.append(_make_segment(buf, chapter_id, scene_id, part))
                buf, size = [], 0
            buf.append(node)
            size += extra
        if buf:
            part += 1
            segments.append(_make_segment(buf, chapter_id, scene_id, part))

    # Files without chapter/scene structure still get deterministic segments.
    if not segments:
        items = [n for n in raw_nodes if n.text]
        buf, size, part = [], 0, 0
        for node in items:
            extra = len(node.text) + (2 if buf else 0)
            if buf and size + extra > max_chars:
                part += 1
                segments.append(_make_segment(buf, None, None, part))
                buf, size = [], 0
            buf.append(node); size += extra
        if buf:
            part += 1
            segments.append(_make_segment(buf, None, None, part))

    return segments


def _make_segment(items: list[DocumentNode], chapter_id: str | None, scene_id: str | None, part: int) -> DocumentSegment:
    first, last = items[0], items[-1]
    chapter_title = None
    if chapter_id:
        # The chapter title is carried as an attribute when available.
        chapter_title = first.attributes.get("chapter_title")
    title = chapter_title or first.attributes.get("section_title") or (first.text[:100] if first.text else "segment")
    lines = []
    for n in items:
        # Include a stable node marker in the LLM input so evidence can be returned precisely.
        lines.append(f"[{n.id}] {n.text}")
    locator = _merge_locators(first.locator, last.locator)
    return DocumentSegment(
        id=f"segment:{_slug(chapter_id or 'document')}:{_slug(scene_id or 'all')}:{part}",
        title=title,
        node_ids=[n.id for n in items],
        section_ids=list(dict.fromkeys(x for x in [chapter_id, scene_id] if x)),
        text="\n\n".join(lines),
        locator=locator,
        chapter_id=chapter_id,
        scene_id=scene_id,
    )


def _merge_locators(a: dict[str, Any], b: dict[str, Any]) -> dict[str, Any]:
    out = dict(a or {})
    for key, value in (b or {}).items():
        if key not in out:
            out[key] = value
        elif out[key] != value:
            if key.endswith("_start"):
                continue
            out[key.replace("_start", "_end")] = value
    return out


def parse_document(path: Path, area: str = "source", max_chars: int = 14000) -> ParsedDocument:
    ext = path.suffix.lower()
    rel = path.as_posix()
    artifact_id = f"{area}:{rel}"
    if ext == ".pdf":
        return _parse_pdf(path, area, rel, artifact_id, max_chars)
    if ext == ".docx":
        return _parse_docx(path, area, rel, artifact_id, max_chars)
    if ext == ".pptx":
        return _parse_pptx(path, area, rel, artifact_id, max_chars)
    if ext == ".xlsx":
        return _parse_xlsx(path, area, rel, artifact_id, max_chars)
    if ext in {".html", ".htm"}:
        return _parse_html(path, area, rel, artifact_id, max_chars)
    if ext in {".md", ".txt", ".csv", ".json", ".yaml", ".yml", ".xml", ".log", ".py", ".js", ".ts", ".css", ".sql"}:
        return _parse_text(path, area, rel, artifact_id, max_chars)
    raise ValueError(f"Unsupported structured document format: {ext}")


def _close_structural_ranges(raw_nodes: list[DocumentNode]):
    levels={"chapter":0,"scene":1,"section":2,"subsection":3,"heading":2,"title":0}
    headings=[(i,n) for i,n in enumerate(raw_nodes) if n.type in levels]
    for idx,node in headings:
        level=levels[node.type]
        stop=len(raw_nodes)
        for j,nxt in headings:
            if j<=idx:
                continue
            if levels[nxt.type] <= level:
                stop=j
                break
        candidates=raw_nodes[idx:stop]
        if not candidates:
            continue
        last=next((x for x in reversed(candidates) if x.locator), node)
        loc=dict(node.locator or {})
        end=dict(last.locator or {})
        if "page" in loc or "page_end" in loc:
            loc["page_start"]=loc.get("page_start",loc.get("page"))
            loc["page_end"]=end.get("page_end",end.get("page",loc.get("page_start")))
            if "line" in loc:
                loc["line_start"] = loc.get("line_start", loc.get("line"))
                loc["line_end"] = end.get("line_end", end.get("line", loc.get("line_start")))
        if "paragraph_index" in loc or "paragraph_start" in loc:
            loc["paragraph_start"]=loc.get("paragraph_start",loc.get("paragraph_index"))
            loc["paragraph_end"]=end.get("paragraph_end",end.get("paragraph_index",loc.get("paragraph_start")))
        if "slide" in loc or "slide_start" in loc:
            loc["slide_start"]=loc.get("slide_start",loc.get("slide"))
            loc["slide_end"]=end.get("slide_end",end.get("slide",loc.get("slide_start")))
        if "char_start" in loc:
            loc["char_end"]=end.get("char_end",loc.get("char_end"))
        if "row" in loc or "row_start" in loc:
            loc["row_start"]=loc.get("row_start",loc.get("row"))
            loc["row_end"]=end.get("row_end",end.get("row",loc.get("row_start")))
        node.locator=loc


def _finalize(path, area, rel, artifact_id, fmt, title, nodes, root, raw_nodes, max_chars, metadata=None):
    _close_structural_ranges(raw_nodes)
    # Ensure all content nodes have useful section labels for evidence-aware compilation.
    chapter_titles = {n.id: n.title for n in raw_nodes if n.type == "chapter"}
    scene_titles = {n.id: n.title for n in raw_nodes if n.type == "scene"}
    for n in raw_nodes:
        ch = n.attributes.get("chapter_id")
        sc = n.attributes.get("scene_id")
        if ch:
            n.attributes["chapter_title"] = chapter_titles.get(ch)
        if sc:
            n.attributes["scene_title"] = scene_titles.get(sc)
    segments = _sectionize(nodes, root, raw_nodes, max_chars=max_chars)
    # Persistable node metadata never includes node.text.
    return ParsedDocument(
        artifact_id=artifact_id,
        area=area,
        path=rel,
        format=fmt,
        title=title,
        nodes=nodes,
        root_id=root.id,
        segments=segments,
        metadata=metadata or {},
    )


def _parse_pdf(path, area, rel, artifact_id, max_chars):
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    nodes = {}
    root = _new_node(nodes, "document", 0, title=path.stem, locator={"format": "pdf"})
    raw: list[DocumentNode] = []
    global_offset = 0
    current_chapter = None
    current_scene = None
    current_section = None
    for page_no, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        page_node = _new_node(nodes, "page", page_no, title=f"Page {page_no}", parent_id=root.id,
                              locator={"page": page_no, "page_start": page_no, "page_end": page_no,
                                       "text_offset_start": global_offset, "text_offset_end": global_offset + len(text)},
                              attributes={"page_number": page_no})
        lines = [x.strip() for x in text.splitlines() if x.strip()]
        local = 0
        for line_no, line in enumerate(lines):
            kind, title = _heading_kind(line)
            if kind == "chapter":
                current_chapter = _new_node(nodes, "chapter", page_no * 100000 + line_no, title=title, parent_id=root.id,
                                            locator={"page_start": page_no, "page_end": page_no}, attributes={"chapter_number": _chapter_number(title)})
                current_scene = None; current_section = current_chapter
                raw.append(current_chapter)
                continue
            if kind == "scene":
                current_scene = _new_node(nodes, "scene", page_no * 100000 + line_no, title=title, parent_id=current_chapter.id if current_chapter else root.id,
                                           locator={"page_start": page_no, "page_end": page_no}, attributes={"scene_number": _scene_number(title)})
                current_section = current_scene
                raw.append(current_scene)
                continue
            if kind in {"heading", "section", "subsection", "title"}:
                node = _new_node(nodes, kind, page_no * 100000 + line_no, text=line, title=title,
                                  parent_id=current_section.id if current_section else root.id,
                                  locator={"page": page_no, "line": line_no, "char_start": global_offset + local, "char_end": global_offset + local + len(line)})
                node.attributes.update({"chapter_id": current_chapter.id if current_chapter else None,
                                        "scene_id": current_scene.id if current_scene else None,
                                        "section_id": current_section.id if current_section else None})
                raw.append(node); current_section = node
            else:
                node = _new_node(nodes, "paragraph", page_no * 100000 + line_no, text=line,
                                  parent_id=current_section.id if current_section else root.id,
                                  locator={"page": page_no, "line": line_no, "char_start": global_offset + local, "char_end": global_offset + local + len(line)})
                node.attributes.update({"chapter_id": current_chapter.id if current_chapter else None,
                                        "scene_id": current_scene.id if current_scene else None,
                                        "section_id": current_section.id if current_section else None})
                raw.append(node)
            local += len(line) + 1
        page_node.attributes.update({
            "chapter_id": current_chapter.id if current_chapter else None,
            "scene_id": current_scene.id if current_scene else None,
            "section_id": current_section.id if current_section else None,
        })
        if current_chapter:
            current_chapter.locator["page_end"] = page_no
        if current_scene:
            current_scene.locator["page_end"] = page_no
        global_offset += len(text) + 2
    return _finalize(path, area, rel, artifact_id, "pdf", path.stem, nodes, root, raw, max_chars,
                     {"page_count": len(reader.pages), "locator_kind": "page"})


def _chapter_number(title):
    m = _CHAPTER_RE.match(title or "")
    return m.group(1) if m else None

def _scene_number(title):
    m = _SCENE_RE.match(title or "")
    return m.group(1) if m else None


def _parse_docx(path, area, rel, artifact_id, max_chars):
    from docx import Document
    doc = Document(str(path))
    nodes = {}; root = _new_node(nodes, "document", 0, title=path.stem, locator={"format":"docx"}); raw=[]
    current_chapter = current_scene = current_section = None
    order = 1
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if not text: continue
        style = p.style.name if p.style else None
        kind, title = _heading_kind(text, style)
        loc = {"paragraph_index": i, "paragraph_start": i, "paragraph_end": i, "style": style}
        if kind == "chapter":
            current_chapter = _new_node(nodes, "chapter", order, title=title, parent_id=root.id, locator=loc, attributes={"chapter_number": _chapter_number(title)}); order+=1
            current_scene=None; current_section=current_chapter; raw.append(current_chapter); continue
        if kind == "scene":
            current_scene = _new_node(nodes, "scene", order, title=title, parent_id=current_chapter.id if current_chapter else root.id, locator=loc, attributes={"scene_number": _scene_number(title)}); order+=1
            current_section=current_scene; raw.append(current_scene); continue
        node = _new_node(nodes, kind or "paragraph", order, text=text, title=title,
                         parent_id=current_section.id if current_section else root.id, locator=loc,
                         attributes={"chapter_id": current_chapter.id if current_chapter else None,
                                     "scene_id": current_scene.id if current_scene else None,
                                     "section_id": current_section.id if current_section else None,
                                     "style": style,
                                     "runs": [{"text":r.text,"bold":r.bold,"italic":r.italic,"underline":r.underline} for r in p.runs]})
        raw.append(node); order += 1
    for ti, table in enumerate(doc.tables):
        rows=[]
        for row in table.rows:
            rows.append([cell.text for cell in row.cells])
        node=_new_node(nodes,"table",100000+ti,text="\n".join(" | ".join(r) for r in rows),parent_id=current_section.id if current_section else root.id,
                       locator={"table_index":ti},attributes={"rows":len(rows),"columns":max((len(r) for r in rows),default=0)})
        raw.append(node)
    return _finalize(path, area, rel, artifact_id, "docx", path.stem, nodes, root, raw, max_chars,
                     {"paragraph_count":len(doc.paragraphs),"table_count":len(doc.tables),"locator_kind":"paragraph"})


def _parse_pptx(path, area, rel, artifact_id, max_chars):
    from pptx import Presentation
    prs=Presentation(str(path)); nodes={}; root=_new_node(nodes,"presentation",0,title=path.stem,locator={"format":"pptx"}); raw=[]
    for si, slide in enumerate(prs.slides,1):
        slide_node=_new_node(nodes,"slide",si,title=f"Slide {si}",parent_id=root.id,locator={"slide":si,"slide_start":si,"slide_end":si})
        for shi, shape in enumerate(slide.shapes):
            text=shape.text.strip() if getattr(shape,"has_text_frame",False) else ""
            if not text: continue
            kind="shape"
            title=None
            if getattr(shape,"is_placeholder",False):
                try:
                    ptype=str(shape.placeholder_format.type)
                    if "TITLE" in ptype.upper(): kind="title"; title=text
                except Exception: pass
            node=_new_node(nodes,kind,si*10000+shi,text=text,title=title,parent_id=slide_node.id,
                           locator={"slide":si,"shape_index":shi,"left":float(shape.left/914400),"top":float(shape.top/914400),"width":float(shape.width/914400),"height":float(shape.height/914400)},
                           attributes={"shape_type":str(shape.shape_type)})
            raw.append(node)
    return _finalize(path, area, rel, artifact_id, "pptx", path.stem, nodes, root, raw, max_chars,
                     {"slide_count":len(prs.slides),"locator_kind":"slide"})


def _parse_xlsx(path, area, rel, artifact_id, max_chars):
    from openpyxl import load_workbook
    wb=load_workbook(path,read_only=True,data_only=False); nodes={}; root=_new_node(nodes,"workbook",0,title=path.stem,locator={"format":"xlsx"}); raw=[]
    for si, ws in enumerate(wb.worksheets):
        sheet=_new_node(nodes,"sheet",si+1,title=ws.title,parent_id=root.id,locator={"sheet":ws.title,"sheet_index":si})
        # One node per non-empty row. This preserves cell coordinates without flattening the workbook.
        for ri,row in enumerate(ws.iter_rows(),1):
            vals=[]; first=None; last=None
            for cell in row:
                if cell.value is not None:
                    vals.append(f"{cell.coordinate}={cell.value}")
                    first=first or cell.coordinate; last=cell.coordinate
            if not vals: continue
            node=_new_node(nodes,"row",si*100000+ri,text=" | ".join(vals),parent_id=sheet.id,
                           locator={"sheet":ws.title,"row":ri,"cell_start":first,"cell_end":last},
                           attributes={"columns":len(vals)})
            raw.append(node)
    return _finalize(path, area, rel, artifact_id, "xlsx", path.stem, nodes, root, raw, max_chars,
                     {"sheet_count":len(wb.worksheets),"locator_kind":"cell"})


def _parse_text(path, area, rel, artifact_id, max_chars):
    text=path.read_text(encoding="utf-8",errors="replace"); nodes={}; root=_new_node(nodes,"document",0,title=path.stem,locator={"format":path.suffix.lower().lstrip(".")}); raw=[]
    current_chapter=current_scene=current_section=None; offset=0; order=1
    blocks=[b.strip() for b in re.split(r"\n\s*\n",text) if b.strip()]
    for block in blocks:
        first=block.splitlines()[0].strip(); kind,title=_heading_kind(first)
        if kind=="chapter":
            current_chapter=_new_node(nodes,"chapter",order,title=title,parent_id=root.id,locator={"char_start":offset,"char_end":offset+len(block)},attributes={"chapter_number":_chapter_number(title)}); order+=1; current_scene=None; current_section=current_chapter; raw.append(current_chapter); offset+=len(block)+2; continue
        if kind=="scene":
            current_scene=_new_node(nodes,"scene",order,title=title,parent_id=current_chapter.id if current_chapter else root.id,locator={"char_start":offset,"char_end":offset+len(block)},attributes={"scene_number":_scene_number(title)}); order+=1; current_section=current_scene; raw.append(current_scene); offset+=len(block)+2; continue
        node=_new_node(nodes,kind or "paragraph",order,text=block,title=title,parent_id=current_section.id if current_section else root.id,
                       locator={"char_start":offset,"char_end":offset+len(block)},attributes={"chapter_id":current_chapter.id if current_chapter else None,"scene_id":current_scene.id if current_scene else None,"section_id":current_section.id if current_section else None})
        raw.append(node); order+=1; offset+=len(block)+2
    return _finalize(path,area,rel,artifact_id,path.suffix.lower().lstrip("."),path.stem,nodes,root,raw,max_chars,{"char_count":len(text),"locator_kind":"char"})


def _parse_html(path, area, rel, artifact_id, max_chars):
    from html.parser import HTMLParser
    class P(HTMLParser):
        def __init__(self): super().__init__(); self.parts=[]; self.stack=[]; self.current=[]
        def handle_starttag(self,tag,attrs): self.stack.append(tag); self.current=[]
        def handle_endtag(self,tag):
            text=" ".join("".join(self.current).split())
            if text and tag in {"p","h1","h2","h3","li","td","th"}: self.parts.append((tag,text))
            if self.stack: self.stack.pop()
        def handle_data(self,data): self.current.append(data)
    p=P(); p.feed(path.read_text(encoding="utf-8",errors="replace")); nodes={}; root=_new_node(nodes,"document",0,title=path.stem,locator={"format":"html"}); raw=[]; current_chapter=None; current_section=None
    for i,(tag,text) in enumerate(p.parts,1):
        kind="paragraph"
        if tag=="h1": kind="chapter"; current_chapter=None
        elif tag=="h2": kind="section"
        elif tag=="h3": kind="subsection"
        parent=current_section.id if current_section else root.id
        if kind=="chapter": current_chapter=_new_node(nodes,"chapter",i,title=text,parent_id=root.id,locator={"dom_order":i},attributes={}); current_section=current_chapter; raw.append(current_chapter); continue
        node=_new_node(nodes,kind,i,text=text,title=text if kind!="paragraph" else None,parent_id=parent,locator={"dom_order":i,"tag":tag},attributes={"chapter_id":current_chapter.id if current_chapter else None}); raw.append(node)
        if kind in {"section","subsection"}: current_section=node
    return _finalize(path,area,rel,artifact_id,"html",path.stem,nodes,root,raw,max_chars,{"locator_kind":"dom_order"})
